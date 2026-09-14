#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
pdf_convert.py — 常用 PDF 互转工具（纯 Python，沙箱友好）
============================================================
支持方向（子命令）：
  excel2pdf    Excel(.xlsx/.xls) -> PDF      （LibreOffice 优先，reportlab 兜底）
  pdf2excel    PDF -> Excel(.xlsx)            （pdfplumber 抽表格 + openpyxl 写出）
  pdf2img      PDF -> 图片(PNG/JPEG)          （pypdfium2 逐页渲染）
  img2pdf      图片(PNG/JPG...) -> PDF        （Pillow 拼成多页 PDF）
  html2pdf     HTML -> PDF                    （LibreOffice 优先，reportlab 兜底）
  pdf2html     PDF -> HTML                    （每页=截图+抽取文本，自包含单文件）
  pdf2docx     PDF -> Word(.docx)             （手写 OOXML：每页=标题+整页图+文本）
  pdf2pptx     PDF -> PPT(.pptx)              （python-pptx：每页=整页图）

设计要点（重要，见 skill 说明）：
  - 沙箱里的 LibreOffice 能写 PDF，但写不出合法的 docx/pptx/xlsx（OOXML 导出失败），
    所以 Word/PPT/Excel 的「写出」一律用纯 Python 库，不依赖 LibreOffice 的 OOXML 导出。
  - LibreOffice 仅在「输入是 Office/HTML、输出是 PDF」时用于提升保真度（沙箱已验证可用）。

用法示例：
  python pdf_convert.py pdf2img  in.pdf --out dir --dpi 150
  python pdf_convert.py img2pdf  a.png b.png --out out.pdf
  python pdf_convert.py pdf2excel in.pdf --out out.xlsx
  python pdf_convert.py excel2pdf data.xlsx --out data.pdf
  python pdf_convert.py html2pdf page.html --out page.pdf
  python pdf_convert.py pdf2html in.pdf --out in.html
  python pdf_convert.py pdf2docx in.pdf --out in.docx
  python pdf_convert.py pdf2pptx in.pdf --out in.pptx
"""
import argparse
import base64
import io
import os
import re
import shutil
import struct
import subprocess
import sys
import tempfile
import zipfile
from pathlib import Path

# ---------- 通用工具 ----------
EMU_PER_IN = 914400
IMG_WIDTH_IN = 6.0


def find_soffice():
    if p := shutil.which("soffice"):
        return p
    for c in [
        r"C:\Program Files\LibreOffice\program\soffice.exe",
        r"C:\Program Files (x86)\LibreOffice\program\soffice.exe",
        os.path.expandvars(r"%LOCALAPPDATA%\LibreOffice\program\soffice.exe"),
        "/Applications/LibreOffice.app/Contents/MacOS/soffice",
        "/usr/bin/soffice",
    ]:
        if os.path.exists(c):
            return c
    return None


def run_lo_convert(src, outdir, fmt, timeout=120):
    """用 LibreOffice 把 src 转成 fmt（如 pdf），成功返回输出文件路径，失败返回 None。"""
    soffice = find_soffice()
    if not soffice:
        return None
    prog = os.path.dirname(soffice)
    profile = tempfile.mkdtemp(prefix="lo_")
    out = Path(outdir)
    out.mkdir(parents=True, exist_ok=True)
    cmd = [
        soffice, "--headless", "--norestore", "--nofirststartwizard", "--nologo",
        f"-env:UserInstallation=file:///{profile.replace(os.sep, '/')}",
        "--convert-to", fmt, "--outdir", str(out), str(src),
    ]
    try:
        subprocess.run(cmd, cwd=prog, capture_output=True, text=True,
                       errors="replace", timeout=timeout)
    except Exception:
        return None
    # 找出产物
    base = Path(src).stem
    cand = out / f"{base}.{fmt}"
    if cand.exists() and cand.stat().st_size > 0:
        return str(cand)
    # 有些版本会带扩展名变体
    for f in out.glob(f"{base}.*"):
        if f.suffix.lower() != ".pdf" or fmt == "pdf":
            if f.stat().st_size > 0:
                return str(f)
    return None


def png_size(path):
    with open(path, "rb") as f:
        f.read(8)
        f.read(4)
        assert f.read(4) == b"IHDR"
        return struct.unpack(">II", f.read(8))


def sanitize_xml_chars(s):
    out = []
    for ch in s:
        o = ord(ch)
        if o in (0x9, 0xA, 0xD) or 0x20 <= o <= 0xD7FF or 0xE000 <= o <= 0xFFFD or 0x10000 <= o <= 0x10FFFF:
            out.append(ch)
    return "".join(out)


def xml_escape(s):
    s = sanitize_xml_chars(s)
    return (s.replace("&", "&amp;").replace("<", "&lt;")
            .replace(">", "&gt;").replace('"', "&quot;"))


# ---------- 1. PDF -> 图片 ----------
def cmd_pdf2img(args):
    import pypdfium2 as pdfium
    out = Path(args.out)
    out.mkdir(parents=True, exist_ok=True)
    doc = pdfium.PdfDocument(args.src)
    scale = args.dpi / 72.0
    fmt = args.format.lower()
    written = []
    pages = range(len(doc)) if args.all else [args.page - 1] if args.page else range(len(doc))
    for i in pages:
        img = doc[i].render(scale=scale).to_pil()
        name = f"{Path(args.src).stem}_page_{i+1:03d}.{fmt}"
        if fmt == "jpeg":
            img = img.convert("RGB")
            img.save(out / name, "JPEG", quality=args.quality)
        else:
            img.save(out / name, "PNG")
        written.append(name)
    print(f"已生成 {len(written)} 张图片 -> {out}")


# ---------- 2. 图片 -> PDF ----------
def cmd_img2pdf(args):
    from PIL import Image
    files = []
    for item in args.images:
        p = Path(item)
        if p.is_dir():
            files += sorted(p.iterdir())
        else:
            files.append(p)
    files = [f for f in files if f.suffix.lower() in (".png", ".jpg", ".jpeg", ".bmp", ".tif", ".tiff")]
    if not files:
        raise SystemExit("没有找到图片文件")
    imgs = [Image.open(f).convert("RGB") for f in files]
    out = args.out
    imgs[0].save(out, "PDF", resolution=100.0, save_all=True, append_images=imgs[1:])
    print(f"已生成: {out} （{len(imgs)} 页）")


# ---------- 3. PDF -> Excel ----------
def sanitize_cell(s):
    """去掉 openpyxl 不允许的控制字符（保留 tab/换行/回车）。"""
    if not isinstance(s, str):
        return s
    return re.sub(r"[\x00-\x08\x0b\x0c\x0e-\x1f]", "", s)


def cmd_pdf2excel(args):
    import pdfplumber
    import openpyxl
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    with pdfplumber.open(args.src) as doc:
        for pi, page in enumerate(doc.pages):
            ws = wb.create_sheet(title=f"Page{pi+1}")
            tables = page.extract_tables()
            row = 1
            if tables:
                for t in tables:
                    for r in t:
                        for c, val in enumerate(r, 1):
                            ws.cell(row=row, column=c, value=sanitize_cell(val or ""))
                        row += 1
                    row += 1
            else:
                txt = page.extract_text() or ""
                for line in txt.split("\n"):
                    ws.cell(row=row, column=1, value=sanitize_cell(line))
                    row += 1
            # 简单列宽
            for col in ws.columns:
                ws.column_dimensions[col[0].column_letter].width = 22
    wb.save(args.out)
    print(f"已生成: {args.out} （{len(wb.sheetnames)} 个工作表）")


# ---------- 4. Excel -> PDF ----------
def cmd_excel2pdf(args):
    # 优先 LibreOffice（保真度高）
    lo = run_lo_convert(args.src, str(Path(args.out).parent), "pdf")
    if lo and os.path.getsize(lo) > 1000:
        # 重命名为期望输出
        if lo != args.out:
            shutil.move(lo, args.out)
        print(f"已生成(LO): {args.out}")
        return
    # 兜底：reportlab 渲染
    print("LibreOffice 不可用或失败，改用 reportlab 兜底渲染。", file=sys.stderr)
    _excel2pdf_reportlab(args.src, args.out)


def _excel2pdf_reportlab(src, out):
    import openpyxl
    from reportlab.lib.pagesizes import landscape, A4
    from reportlab.lib import colors
    from reportlab.lib.units import mm
    from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer
    from reportlab.lib.styles import getSampleStyleSheet

    wb = openpyxl.load_workbook(src, data_only=True)
    doc = SimpleDocTemplate(out, pagesize=landscape(A4))
    style = getSampleStyleSheet()["Normal"]
    style.fontSize = 7
    flow = []
    for ws in wb.worksheets:
        rows = list(ws.iter_rows(values_only=True))
        if not rows:
            continue
        data = [[("" if c is None else str(c)) for c in r] for r in rows]
        ncol = max((len(r) for r in data), default=1)
        data = [r + [""] * (ncol - len(r)) for r in data]
        t = Table(data, repeatRows=1)
        t.setStyle(TableStyle([
            ("FONTSIZE", (0, 0), (-1, -1), 7),
            ("GRID", (0, 0), (-1, -1), 0.25, colors.grey),
            ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#1F4E79")),
            ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
            ("VALIGN", (0, 0), (-1, -1), "TOP"),
            ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#F2F6FB")]),
        ]))
        flow.append(Paragraph(f"<b>{xml_escape(ws.title)}</b>", style))
        flow.append(t)
        flow.append(Spacer(1, 6 * mm))
    doc.build(flow)
    print(f"已生成: {out}")


# ---------- 5. HTML -> PDF ----------
def cmd_html2pdf(args):
    lo = run_lo_convert(args.src, str(Path(args.out).parent), "pdf")
    if lo and os.path.getsize(lo) > 1000:
        if lo != args.out:
            shutil.move(lo, args.out)
        print(f"已生成(LO): {args.out}")
        return
    print("LibreOffice 不可用或失败，改用 reportlab 兜底（仅保留文本）。", file=sys.stderr)
    _html2pdf_reportlab(args.src, args.out)


def _html2pdf_reportlab(src, out):
    from reportlab.lib.pagesizes import A4
    from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer
    from reportlab.lib.styles import getSampleStyleSheet
    from html.parser import HTMLParser

    class T(HTMLParser):
        def __init__(self):
            super().__init__()
            self.parts = []

        def handle_data(self, data):
            if data.strip():
                self.parts.append(data)

    p = T()
    with open(src, encoding="utf-8", errors="replace") as f:
        p.feed(f.read())
    doc = SimpleDocTemplate(out, pagesize=A4)
    ss = getSampleStyleSheet()
    flow = [Paragraph(xml_escape(x), ss["Normal"]) for x in p.parts]
    doc.build(flow)


# ---------- 6. PDF -> HTML ----------
def cmd_pdf2html(args):
    import pdfplumber
    import pypdfium2 as pdfium

    src = args.src
    doc_pdf = pdfium.PdfDocument(src)
    out_parts = []
    with pdfplumber.open(src) as doc:
        for i, page in enumerate(doc.pages):
            # 截图转 base64
            img = doc_pdf[i].render(scale=1.5).to_pil()
            buf = io.BytesIO()
            img.save(buf, "PNG")
            b64 = base64.b64encode(buf.getvalue()).decode()
            text = (page.extract_text() or "").strip()
            paras = "".join(f"<p>{xml_escape(ln.rstrip())}</p>" for ln in text.split("\n") if ln.strip())
            out_parts.append(
                f'<section class="pdf-page">'
                f'<h2>第 {i+1} 页</h2>'
                f'<img src="data:image/png;base64,{b64}" alt="page {i+1}"/>'
                f'<div class="text">{paras}</div>'
                f'</section>'
            )
    html = (
        '<!DOCTYPE html><html lang="zh"><head><meta charset="utf-8">'
        '<title>PDF 导出</title>'
        '<style>body{font-family:system-ui,"Microsoft YaHei",sans-serif;background:#f5f5f5;'
        'margin:0;padding:24px;}'
        '.pdf-page{background:#fff;max-width:900px;margin:0 auto 24px;padding:24px;'
        'box-shadow:0 1px 4px rgba(0,0,0,.15);border-radius:6px;}'
        '.pdf-page img{max-width:100%;height:auto;border:1px solid #eee;}'
        '.text{margin-top:16px;white-space:pre-wrap;line-height:1.6;color:#222;}</style>'
        f'</head><body>{"".join(out_parts)}</body></html>'
    )
    Path(args.out).write_text(html, encoding="utf-8")
    print(f"已生成: {args.out}")


# ---------- 7. PDF -> DOCX（手写 OOXML）----------
STYLES_XML = (
    '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
    '<w:styles xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">'
    '<w:docDefaults><w:rPrDefault><w:rPr>'
    '<w:rFonts w:ascii="Calibri" w:hAnsi="Calibri" w:eastAsia="SimSun"/>'
    '<w:sz w:val="21"/></w:rPr></w:rPrDefault></w:docDefaults>'
    '<w:style w:type="paragraph" w:default="1" w:styleId="Normal"><w:name w:val="Normal"/></w:style>'
    '<w:style w:type="paragraph" w:styleId="Heading1"><w:name w:val="heading 1"/>'
    '<w:basedOn w:val="Normal"/><w:pPr><w:keepNext/><w:spacing w:before="240" w:after="80"/>'
    '<w:outlineLvl w:val="0"/></w:pPr>'
    '<w:rPr><w:b/><w:sz w:val="32"/><w:color w:val="1F4E79"/></w:rPr></w:style>'
    '</w:styles>'
)
CORE_XML = (
    '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
    '<cp:coreProperties xmlns:cp="http://schemas.openxmlformats.org/package/2006/metadata/core-properties" '
    'xmlns:dc="http://purl.org/dc/elements/1.1/" xmlns:dcterms="http://purl.org/dc/terms/" '
    'xmlns:xsi="http://www.w3.org/2001/XMLSchema-instance">'
    '<dc:title>PDF 转换文档</dc:title><dc:creator>WorkBuddy</dc:creator></cp:coreProperties>'
)
APP_XML = (
    '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
    '<Properties xmlns="http://schemas.openxmlformats.org/officeDocument/2006/extended-properties">'
    '<Application>WorkBuddy PDF2Docx</Application></Properties>'
)


def cmd_pdf2docx(args):
    import pdfplumber
    import pypdfium2 as pdfium

    # 渲染整页图
    rendered = []
    doc = pdfium.PdfDocument(args.src)
    tmp = Path(tempfile.mkdtemp(prefix="docxpg_"))
    for i in range(len(doc)):
        img = doc[i].render(scale=2.0).to_pil()
        fp = tmp / f"page_{i+1:03d}.png"
        img.save(fp)
        rendered.append(str(fp))

    with pdfplumber.open(args.src) as pdf:
        texts = [(p.extract_text() or "").strip() for p in pdf.pages]

    n = len(texts)
    body, rels = [], []
    rid = 100
    for i in range(n):
        body.append(f'<w:p><w:pPr><w:pStyle w:val="Heading1"/></w:pPr>'
                    f'<w:r><w:t xml:space="preserve">第 {i+1} 页</w:t></w:r></w:p>')
        wpx, hpx = png_size(rendered[i])
        cx = int(IMG_WIDTH_IN * EMU_PER_IN)
        cy = int(cx * hpx / wpx)
        r = f"rId{rid}"; rid += 1
        rels.append((r, f"media/image{i+1}.png"))
        body.append(
            f'<w:p><w:r><w:drawing><wp:inline '
            f'xmlns:wp="http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing" '
            f'distT="0" distB="0" distL="0" distR="0">'
            f'<wp:extent cx="{cx}" cy="{cy}"/>'
            f'<wp:effectExtent l="0" t="0" r="0" b="0"/>'
            f'<wp:docPr id="{i+1}" name="Picture{i+1}"/>'
            f'<wp:cNvGraphicFramePr><a:graphicFrameLocks '
            f'xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main" noChangeAspect="1"/>'
            f'</wp:cNvGraphicFramePr>'
            f'<a:graphic xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main">'
            f'<a:graphicData uri="http://schemas.openxmlformats.org/drawingml/2006/picture">'
            f'<pic:pic xmlns:pic="http://schemas.openxmlformats.org/drawingml/2006/picture">'
            f'<pic:nvPicPr><pic:cNvPr id="{i+1}" name="img{i+1}.png"/>'
            f'<pic:cNvPicPr/></pic:nvPicPr>'
            f'<pic:blipFill><a:blip r:embed="{r}"/><a:stretch><a:fillRect/></a:stretch></pic:blipFill>'
            f'<pic:spPr><a:xfrm><a:off x="0" y="0"/><a:ext cx="{cx}" cy="{cy}"/></a:xfrm>'
            f'<a:prstGeom prst="rect"><a:avLst/></a:prstGeom></pic:spPr>'
            f'</pic:pic></a:graphicData></a:graphic></wp:inline></w:drawing></w:r></w:p>'
        )
        txt = texts[i]
        if txt:
            for line in txt.split("\n"):
                line = line.rstrip()
                if not line.strip():
                    body.append("<w:p/>")
                else:
                    body.append(f'<w:p><w:r><w:t xml:space="preserve">{xml_escape(line)}</w:t></w:r></w:p>')
        else:
            body.append('<w:p><w:r><w:t xml:space="preserve">（本页无可提取文字，见上方截图）</w:t></w:r></w:p>')

    document_xml = (
        '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>\n'
        '<w:document xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main" '
        'xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships">'
        '<w:body>' + "".join(body) +
        '<w:sectPr><w:pgSz w:w="11906" w:h="16838"/>'
        '<w:pgMar w:top="1440" w:right="1440" w:bottom="1440" w:left="1440" '
        'w:header="720" w:footer="720" w:gutter="0"/></w:sectPr>'
        '</w:body></w:document>'
    )
    doc_rels = ['<?xml version="1.0" encoding="UTF-8" standalone="yes"?>',
                '<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">',
                '<Relationship Id="rId1" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/styles" Target="styles.xml"/>']
    for r, target in rels:
        doc_rels.append(f'<Relationship Id="{r}" '
                         f'Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/image" '
                         f'Target="{target}"/>')
    doc_rels.append('</Relationships>')

    with zipfile.ZipFile(args.out, "w", zipfile.ZIP_DEFLATED) as z:
        z.writestr("[Content_Types].xml",
                   '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
                   '<Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types">'
                   '<Default Extension="rels" ContentType="application/vnd.openxmlformats-package.relationships+xml"/>'
                   '<Default Extension="xml" ContentType="application/xml"/>'
                   '<Default Extension="png" ContentType="image/png"/>'
                   '<Override PartName="/word/document.xml" ContentType="application/vnd.openxmlformats-officedocument.wordprocessingml.document.main+xml"/>'
                   '<Override PartName="/word/styles.xml" ContentType="application/vnd.openxmlformats-officedocument.wordprocessingml.styles+xml"/>'
                   '<Override PartName="/docProps/core.xml" ContentType="application/vnd.openxmlformats-package.core-properties+xml"/>'
                   '<Override PartName="/docProps/app.xml" ContentType="application/vnd.openxmlformats-officedocument.extended-properties+xml"/>'
                   '</Types>')
        z.writestr("_rels/.rels",
                   '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
                   '<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">'
                   '<Relationship Id="rId1" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/officeDocument" Target="word/document.xml"/>'
                   '<Relationship Id="rId2" Type="http://schemas.openxmlformats.org/package/2006/relationships/metadata/core-properties" Target="docProps/core.xml"/>'
                   '<Relationship Id="rId3" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/extended-properties" Target="docProps/app.xml"/>'
                   '</Relationships>')
        z.writestr("word/document.xml", document_xml)
        z.writestr("word/_rels/document.xml.rels", "".join(doc_rels))
        z.writestr("word/styles.xml", STYLES_XML)
        z.writestr("docProps/core.xml", CORE_XML)
        z.writestr("docProps/app.xml", APP_XML)
        for i, (r, target) in enumerate(rels):
            z.write(rendered[i], f"word/{target}")
    print(f"已生成: {args.out} （共 {n} 页，{len(rels)} 张图）")


# ---------- 8. PDF -> PPTX ----------
def cmd_pdf2pptx(args):
    from pptx import Presentation
    from pptx.util import Emu
    from PIL import Image

    src = args.src
    doc = pypdfium2_import().PdfDocument(src)
    tmp = Path(tempfile.mkdtemp(prefix="pptxpg_"))
    files = []
    for i in range(len(doc)):
        img = doc[i].render(scale=2.0).to_pil()
        fp = tmp / f"page_{i+1:03d}.png"
        img.save(fp)
        files.append(fp)

    prs = Presentation()
    prs.slides._sldIdLst.clear()
    with Image.open(files[0]) as im:
        w, h = im.size
    aspect = w / h
    base = Emu(12192000)
    if aspect >= 1:
        prs.slide_width = base
        prs.slide_height = Emu(int(base / aspect))
    else:
        prs.slide_height = Emu(9144000)
        prs.slide_width = Emu(int(prs.slide_height * aspect))
    blank = prs.slide_layouts[6]
    for fp in files:
        slide = prs.slides.add_slide(blank)
        slide.shapes.add_picture(str(fp), 0, 0, width=prs.slide_width, height=prs.slide_height)
    prs.save(args.out)
    print(f"已生成: {args.out} （共 {len(files)} 页）")


def pypdfium2_import():
    import pypdfium2
    return pypdfium2


# ---------- CLI ----------
def build_parser():
    p = argparse.ArgumentParser(description="PDF 常用格式互转工具（纯 Python）")
    sub = p.add_subparsers(dest="cmd", required=True)

    s = sub.add_parser("pdf2img"); s.add_argument("src"); s.add_argument("--out", required=True)
    s.add_argument("--dpi", type=int, default=150); s.add_argument("--format", default="png")
    s.add_argument("--quality", type=int, default=90); s.add_argument("--page", type=int, default=None)
    s.add_argument("--all", action="store_true"); s.set_defaults(func=cmd_pdf2img)

    s = sub.add_parser("img2pdf"); s.add_argument("images", nargs="+"); s.add_argument("--out", required=True)
    s.set_defaults(func=cmd_img2pdf)

    s = sub.add_parser("pdf2excel"); s.add_argument("src"); s.add_argument("--out", required=True)
    s.set_defaults(func=cmd_pdf2excel)

    s = sub.add_parser("excel2pdf"); s.add_argument("src"); s.add_argument("--out", required=True)
    s.set_defaults(func=cmd_excel2pdf)

    s = sub.add_parser("html2pdf"); s.add_argument("src"); s.add_argument("--out", required=True)
    s.set_defaults(func=cmd_html2pdf)

    s = sub.add_parser("pdf2html"); s.add_argument("src"); s.add_argument("--out", required=True)
    s.set_defaults(func=cmd_pdf2html)

    s = sub.add_parser("pdf2docx"); s.add_argument("src"); s.add_argument("--out", required=True)
    s.set_defaults(func=cmd_pdf2docx)

    s = sub.add_parser("pdf2pptx"); s.add_argument("src"); s.add_argument("--out", required=True)
    s.set_defaults(func=cmd_pdf2pptx)
    return p


def main():
    args = build_parser().parse_args()
    args.func(args)


if __name__ == "__main__":
    main()
